import hashlib
import json
import logging
import tempfile
import uuid
from decimal import Decimal
from pathlib import Path
from unittest.mock import MagicMock, patch

from django.conf import settings
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import Client, RequestFactory, TestCase
from django.urls import reverse

from core.db import configure_database_connection
from core.file_uploader import FileUploader
from core.logging import InterceptHandler, setup_logging
from core.models import InvestigationProfile
from core.modules import get_discovered_modules


class CoreFileUploaderTests(TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.uploader = FileUploader(self.temp_dir.name, default_ext=".bin")

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_single_chunk_upload_and_sha256(self):
        data = b"FORENSIC_EVIDENCE_PAYLOAD_CHUNK_12345"
        expected_sha = hashlib.sha256(data).hexdigest()

        res = self.uploader.append_chunk(
            upload_id="test_upload_01",
            chunk_index=0,
            total_chunks=1,
            chunk_data=data,
        )

        self.assertTrue(res["is_completed"])
        self.assertEqual(res["file_sha256"], expected_sha)
        self.assertEqual(res["current_size_bytes"], len(data))
        self.assertTrue(Path(res["file_path"]).exists())

    def test_multi_chunk_assembly(self):
        chunk1 = b"PART_ONE_"
        chunk2 = b"PART_TWO_"
        chunk3 = b"PART_THREE"
        full_data = chunk1 + chunk2 + chunk3
        expected_sha = hashlib.sha256(full_data).hexdigest()

        res1 = self.uploader.append_chunk(
            upload_id="multi_01", chunk_index=0, total_chunks=3, chunk_data=chunk1
        )
        self.assertFalse(res1["is_completed"])

        res2 = self.uploader.append_chunk(
            upload_id="multi_01", chunk_index=1, total_chunks=3, chunk_data=chunk2
        )
        self.assertFalse(res2["is_completed"])

        res3 = self.uploader.append_chunk(
            upload_id="multi_01", chunk_index=2, total_chunks=3, chunk_data=chunk3
        )
        self.assertTrue(res3["is_completed"])
        self.assertEqual(res3["file_sha256"], expected_sha)
        self.assertEqual(res3["current_size_bytes"], len(full_data))

    def test_extension_normalization_and_chunk0_overwrite(self):
        res1 = self.uploader.append_chunk(
            upload_id="no_dot_ext",
            chunk_index=0,
            total_chunks=1,
            chunk_data=b"Initial",
            extension="dat",
        )
        self.assertTrue(res1["is_completed"])
        target = Path(res1["file_path"])
        self.assertTrue(target.exists())
        self.assertEqual(target.suffix, ".dat")

        res2 = self.uploader.append_chunk(
            upload_id="no_dot_ext",
            chunk_index=0,
            total_chunks=1,
            chunk_data=b"Overwritten",
            extension=".dat",
        )
        self.assertTrue(res2["is_completed"])
        self.assertEqual(target.read_bytes(), b"Overwritten")


class CorePortalAuthMiddlewareAndViewsTests(TestCase):
    def setUp(self):
        self.client = Client()

    def test_unauthenticated_redirect_to_login(self):
        res = self.client.get("/")
        self.assertEqual(res.status_code, 302)
        self.assertIn("/login/?next=/", res.url)

    def test_exempt_path_permits_anonymous_access(self):
        res = self.client.get("/login/")
        self.assertEqual(res.status_code, 200)
        self.assertIn(b"Authorized Personnel Only", res.content)

    def test_login_failure_with_wrong_password(self):
        res = self.client.post("/login/", {"password": "wrong_password", "next": "/"})
        self.assertEqual(res.status_code, 200)
        self.assertIn(b"Invalid portal access key", res.content)
        self.assertFalse(self.client.session.get("portal_authenticated", False))

    def test_login_success_with_valid_password(self):
        valid_pwd = getattr(settings, "PORTAL_ACCESS_PASSWORD", "forensiq2026")
        res = self.client.post("/login/", {"password": valid_pwd, "next": "/bank/"})
        self.assertEqual(res.status_code, 302)
        self.assertEqual(res.url, "/bank/")
        self.assertTrue(self.client.session.get("portal_authenticated", False))

    def test_open_redirect_sanitization(self):
        valid_pwd = getattr(settings, "PORTAL_ACCESS_PASSWORD", "forensiq2026")
        # Attempt malicious off-site redirect
        res = self.client.post(
            "/login/", {"password": valid_pwd, "next": "https://malicious-site.com"}
        )
        self.assertEqual(res.status_code, 302)
        self.assertEqual(res.url, "/")

        res_double_slash = self.client.post(
            "/login/", {"password": valid_pwd, "next": "//malicious-site.com"}
        )
        self.assertEqual(res_double_slash.status_code, 302)
        self.assertEqual(res_double_slash.url, "/")

    def test_portal_logout_locks_workstation(self):
        # Authenticate first
        session = self.client.session
        session["portal_authenticated"] = True
        session.save()

        # Logout
        res = self.client.get(reverse("portal_logout"))
        self.assertEqual(res.status_code, 302)
        self.assertEqual(res.url, "/login/")
        self.assertFalse(self.client.session.get("portal_authenticated", False))

    def test_landing_view_authenticated(self):
        session = self.client.session
        session["portal_authenticated"] = True
        session.save()

        res = self.client.get("/")
        self.assertEqual(res.status_code, 200)
        self.assertIn(b"Investigation Platform", res.content)


class CoreModulesDiscoveryTests(TestCase):
    def test_get_discovered_modules_structure(self):
        modules = get_discovered_modules()
        self.assertIsInstance(modules, list)
        self.assertGreaterEqual(len(modules), 8)

        app_names = [m["app_name"] for m in modules]
        self.assertIn("q_bank", app_names)
        self.assertIn("q_mail", app_names)
        self.assertIn("q_voice", app_names)
        self.assertNotIn("q_timeline", app_names)

        # Check expected keys in each module
        for m in modules:
            self.assertIn("app_name", m)
            self.assertIn("name", m)
            self.assertIn("tag", m)
            self.assertIn("accent", m)
            self.assertIn("href", m)
            self.assertIn("tagline", m)

    def test_missing_apps_dir(self):
        with patch("core.modules.settings.BASE_DIR", Path("/non_existent_folder_xyz")):
            modules = get_discovered_modules()
            self.assertEqual(modules, [])

    def test_discovered_modules_fallback_attributes(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            apps_dir = temp_path / "apps"
            apps_dir.mkdir()

            # 1. Module without AppConfig or default spec
            custom_app = apps_dir / "q_custom_inspect"
            custom_app.mkdir()
            (custom_app / "__init__.py").write_text("", encoding="utf-8")

            # 2. Excluded module
            excluded_app = apps_dir / "q_timeline"
            excluded_app.mkdir()
            (excluded_app / "__init__.py").write_text("", encoding="utf-8")

            with patch("core.modules.settings.BASE_DIR", temp_path):
                modules = get_discovered_modules()
                self.assertEqual(len(modules), 1)
                mod = modules[0]
                self.assertEqual(mod["app_name"], "q_custom_inspect")
                self.assertEqual(mod["num"], "01")
                self.assertEqual(mod["name"], "Custom Inspect")


class CoreLoggingAndDatabaseTests(TestCase):
    def test_intercept_handler_emit(self):
        handler = InterceptHandler()
        # Normal log record
        record_info = logging.LogRecord(
            name="test_logger",
            level=logging.INFO,
            pathname=__file__,
            lineno=10,
            msg="Forensic audit test message",
            args=(),
            exc_info=None,
        )
        handler.emit(record_info)

        # Custom/unknown level record triggering ValueError in logger.level
        record_custom = logging.LogRecord(
            name="test_logger",
            level=99,
            pathname=__file__,
            lineno=20,
            msg="Custom level test message",
            args=(),
            exc_info=None,
        )
        record_custom.levelname = "CUSTOM_LEVEL_UNKNOWN_99"
        handler.emit(record_custom)

    def test_setup_logging_initialization(self):
        # Call setup_logging and verify it completes without error
        setup_logging()

    def test_configure_database_connection(self):
        # 1. Non-sqlite connection is skipped
        mock_pg_conn = MagicMock()
        mock_pg_conn.vendor = "postgresql"
        configure_database_connection(sender=None, connection=mock_pg_conn)
        mock_pg_conn.cursor.assert_not_called()

        # 2. SQLite connection executes pragmas
        mock_sqlite_conn = MagicMock()
        mock_sqlite_conn.vendor = "sqlite"
        mock_cursor = MagicMock()
        mock_sqlite_conn.cursor.return_value.__enter__.return_value = mock_cursor
        configure_database_connection(sender=None, connection=mock_sqlite_conn)
        self.assertGreaterEqual(mock_cursor.execute.call_count, 4)

        # 3. SQLite connection cursor exception handled gracefully
        mock_err_conn = MagicMock()
        mock_err_conn.vendor = "sqlite"
        mock_err_conn.cursor.side_effect = RuntimeError("Pragma lock error")
        configure_database_connection(sender=None, connection=mock_err_conn)


class CoreInvestigationProfilesTests(TestCase):
    def setUp(self):
        self.factory = RequestFactory()
        self.profile = InvestigationProfile.objects.create(
            full_name="Target Custodian A",
            employee_id="EMP-1001",
            department="Procurement",
            designation="Manager",
            email="custodian.a@example.com",
            phone="+91 9876543210",
            risk_level="HIGH",
            status="ACTIVE",
            notes="Under observation",
            avatar_color="orange",
        )

    def test_get_all_profiles(self):
        from core.profiles import get_all_profiles

        profiles = get_all_profiles()
        self.assertGreaterEqual(len(profiles), 1)
        self.assertEqual(profiles.first().full_name, "Target Custodian A")

    def test_get_profile_by_id(self):
        from core.profiles import get_profile_by_id

        self.assertIsNone(get_profile_by_id(None))
        self.assertIsNone(get_profile_by_id("invalid-uuid"))
        found = get_profile_by_id(self.profile.id)
        self.assertIsNotNone(found)
        self.assertEqual(found.id, self.profile.id)

    def test_get_and_set_active_profile(self):
        from core.profiles import get_active_profile, set_active_profile

        request = self.factory.get("/")
        request.session = {}

        # Initially none
        self.assertIsNone(get_active_profile(request))

        # Set active
        active = set_active_profile(request, self.profile.id)
        self.assertIsNotNone(active)
        self.assertEqual(request.session.get("active_profile_id"), str(self.profile.id))
        self.assertEqual(get_active_profile(request).id, self.profile.id)

        # Clear active
        cleared = set_active_profile(request, None)
        self.assertIsNone(cleared)
        self.assertNotIn("active_profile_id", request.session)

    def test_create_investigation_profile(self):
        from core.profiles import create_investigation_profile

        with self.assertRaises(ValueError):
            create_investigation_profile(full_name="   ")

        p = create_investigation_profile(
            full_name="New Auditee B",
            employee_id="EMP-2002",
            department="Finance",
            risk_level="INVALID_RISK",  # invalid choice falls back to MEDIUM
            status="INVALID",  # invalid choice falls back to ACTIVE
        )
        self.assertEqual(p.full_name, "New Auditee B")
        self.assertEqual(p.risk_level, "MEDIUM")
        self.assertEqual(p.status, "ACTIVE")

    def test_resolve_or_create_profile_from_request(self):
        from core.profiles import resolve_or_create_profile_from_request, set_active_profile

        # 1. Existing Profile Selected
        req1 = self.factory.post("/", {"profile_id": str(self.profile.id)})
        req1.session = {}
        prof, name = resolve_or_create_profile_from_request(req1)
        self.assertEqual(prof.id, self.profile.id)
        self.assertEqual(name, self.profile.full_name)

        # 2. Inline New Profile
        req2 = self.factory.post(
            "/",
            {
                "new_profile_name": "Inline Person",
                "new_profile_dept": "IT",
                "new_profile_role": "Admin",
            },
        )
        req2.session = {}
        prof2, name2 = resolve_or_create_profile_from_request(req2)
        self.assertEqual(prof2.full_name, "Inline Person")
        self.assertEqual(name2, "Inline Person")

        # 2b. Inline Profile that already exists
        req2b = self.factory.post("/", {"new_profile_name": self.profile.full_name})
        req2b.session = {}
        prof2b, name2b = resolve_or_create_profile_from_request(req2b)
        self.assertEqual(prof2b.id, self.profile.id)

        # 3. Legacy Name Fallback
        req3 = self.factory.post("/", {"custodian_name": "Legacy Target"})
        req3.session = {}
        prof3, name3 = resolve_or_create_profile_from_request(req3, default_department="Logistics")
        self.assertEqual(prof3.full_name, "Legacy Target")
        self.assertEqual(prof3.department, "Logistics")

        # 3b. Legacy Name that already exists
        req3b = self.factory.post("/", {"account_holder": self.profile.full_name})
        req3b.session = {}
        prof3b, name3b = resolve_or_create_profile_from_request(req3b)
        self.assertEqual(prof3b.id, self.profile.id)

        # 4. Fallback to active session profile
        req4 = self.factory.post("/", {})
        req4.session = {}
        set_active_profile(req4, self.profile.id)
        prof4, name4 = resolve_or_create_profile_from_request(req4)
        self.assertEqual(prof4.id, self.profile.id)

        # 5. Empty
        req5 = self.factory.post("/", {})
        req5.session = {}
        prof5, name5 = resolve_or_create_profile_from_request(req5)
        self.assertIsNone(prof5)
        self.assertEqual(name5, "")

    def test_sync_all_existing_entities_to_profiles(self):
        from django.utils import timezone
        from q_bank.models import AuditedPerson
        from q_verify.models import VerificationCase
        from q_voice.models import AudioRecording

        from core.profiles import sync_all_existing_entities_to_profiles

        AuditedPerson.objects.create(full_name="Synced Bank Auditee", notes="flagged transaction")
        AudioRecording.objects.create(
            call_ref="CALL-SYNC-001",
            call_title="Synchronized Wiretap",
            call_timestamp=timezone.now(),
            custodian_name="Synced Voice Subject",
            risk_score=85,
        )
        VerificationCase.objects.create(
            custodian_name="Synced Verify Subject", custodian_department="HR"
        )

        created = sync_all_existing_entities_to_profiles()
        self.assertGreaterEqual(created, 3)


class CoreProfileViewsTests(TestCase):
    def setUp(self):
        self.client = Client()
        session = self.client.session
        session["portal_authenticated"] = True
        session.save()

        self.profile = InvestigationProfile.objects.create(
            full_name="Target Person X",
            employee_id="EMP-9009",
            department="Operations",
        )

    def test_create_profile_view_json(self):
        res = self.client.post(
            reverse("create_profile"),
            data=json.dumps({"full_name": "AJAX Created Person", "department": "Legal"}),
            content_type="application/json",
        )
        self.assertEqual(res.status_code, 200)
        data = res.json()
        self.assertEqual(data["status"], "success")
        self.assertEqual(data["profile"]["full_name"], "AJAX Created Person")

    def test_create_profile_view_json_empty_name(self):
        res = self.client.post(
            reverse("create_profile"),
            data=json.dumps({"full_name": "  "}),
            content_type="application/json",
        )
        self.assertEqual(res.status_code, 400)
        data = res.json()
        self.assertEqual(data["status"], "error")

    def test_create_profile_view_form(self):
        res = self.client.post(
            reverse("create_profile"),
            data={"full_name": "Form Created Person", "department": "Audit"},
        )
        self.assertEqual(res.status_code, 302)
        self.assertTrue(
            InvestigationProfile.objects.filter(full_name="Form Created Person").exists()
        )

    def test_set_active_profile_view_json(self):
        # Set active
        res1 = self.client.post(
            reverse("set_active_profile"),
            data=json.dumps({"profile_id": str(self.profile.id)}),
            content_type="application/json",
        )
        self.assertEqual(res1.status_code, 200)
        self.assertEqual(res1.json()["active_profile"]["id"], str(self.profile.id))

        # Clear active
        res2 = self.client.post(
            reverse("set_active_profile"),
            data=json.dumps({"profile_id": ""}),
            content_type="application/json",
        )
        self.assertEqual(res2.status_code, 200)
        self.assertIsNone(res2.json()["active_profile"])

    def test_create_profile_view_form_empty(self):
        res = self.client.post(
            reverse("create_profile"),
            data={"full_name": ""},
        )
        self.assertEqual(res.status_code, 302)

    def test_set_active_profile_view_form(self):
        res = self.client.post(
            reverse("set_active_profile"),
            data={"profile_id": str(self.profile.id)},
        )
        self.assertEqual(res.status_code, 302)

    def test_profile_list_api_view(self):
        res = self.client.get(reverse("api_profiles"))
        self.assertEqual(res.status_code, 200)
        data = res.json()
        self.assertEqual(data["status"], "success")
        self.assertGreaterEqual(len(data["profiles"]), 1)

    def test_create_profile_view_malformed_json(self):
        res = self.client.post(
            reverse("create_profile"),
            data=b"invalid-json{",
            content_type="application/json",
        )
        self.assertEqual(res.status_code, 400)

    def test_set_active_profile_view_malformed_json(self):
        res = self.client.post(
            reverse("set_active_profile"),
            data=b"invalid-json{",
            content_type="application/json",
        )
        self.assertEqual(res.status_code, 200)

    def test_set_active_profile_invalid_id(self):
        from core.profiles import set_active_profile

        request = RequestFactory().post("/")
        request.session = {}
        result = set_active_profile(request, "00000000-0000-0000-0000-000000000000")
        self.assertIsNone(result)

    def test_context_processor_error_fallback(self):
        from core.context_processors import global_profiles_context

        request = RequestFactory().get("/")
        request.session = {}
        with patch(
            "core.context_processors.get_all_profiles", side_effect=RuntimeError("DB offline")
        ):
            ctx = global_profiles_context(request)
            self.assertEqual(ctx["investigation_profiles"], [])
            self.assertEqual(ctx["total_profiles_count"], 0)

    def test_create_profile_with_keywords(self):
        from core.profiles import create_investigation_profile

        prof = create_investigation_profile(
            full_name="Keyword Subject",
            department="Risk & Audit",
            keywords=["Kickback", "commission", "KICKBACK", "off-book"],
        )
        self.assertEqual(prof.keywords, ["Kickback", "commission", "off-book"])
        self.assertIn("keywords", prof.to_dict())

    def test_add_keywords_to_profile(self):
        from core.profiles import add_keywords_to_profile, create_investigation_profile

        prof = create_investigation_profile(
            full_name="Surveillance Target",
            keywords=["bribe"],
        )
        updated = add_keywords_to_profile(prof.id, "cash, gift, bribe, secret")
        self.assertEqual(updated.keywords, ["bribe", "cash", "gift", "secret"])

    def test_add_keywords_to_profile_not_found(self):
        from core.profiles import add_keywords_to_profile

        with self.assertRaises(ValueError):
            add_keywords_to_profile(uuid.uuid4(), "audit")

    def test_create_profile_view_with_keywords_json(self):
        res = self.client.post(
            reverse("create_profile"),
            data=json.dumps(
                {
                    "full_name": "Profile With Keywords",
                    "department": "Security",
                    "keywords": ["shell_company", "wire_transfer"],
                }
            ),
            content_type="application/json",
        )
        self.assertEqual(res.status_code, 200)
        data = res.json()
        self.assertEqual(data["profile"]["keywords"], ["shell_company", "wire_transfer"])

    def test_create_profile_view_with_keywords_form(self):
        res = self.client.post(
            reverse("create_profile"),
            data={
                "full_name": "Form Profile Keywords",
                "department": "Procurement",
                "keywords": json.dumps(["cash_payment", "hawala"]),
            },
        )
        self.assertEqual(res.status_code, 302)
        prof = InvestigationProfile.objects.filter(full_name="Form Profile Keywords").first()
        self.assertIsNotNone(prof)
        self.assertEqual(prof.keywords, ["cash_payment", "hawala"])

    def test_add_profile_keywords_view_success(self):
        url = reverse("add_profile_keywords", kwargs={"profile_id": str(self.profile.id)})
        res = self.client.post(
            url,
            data=json.dumps({"keywords": "fraud, evasion"}),
            content_type="application/json",
        )
        self.assertEqual(res.status_code, 200)
        data = res.json()
        self.assertEqual(data["status"], "success")
        self.assertIn("fraud", data["profile"]["keywords"])
        self.assertIn("evasion", data["profile"]["keywords"])

    def test_add_profile_keywords_view_empty(self):
        url = reverse("add_profile_keywords", kwargs={"profile_id": str(self.profile.id)})
        res = self.client.post(
            url,
            data=json.dumps({"keywords": ""}),
            content_type="application/json",
        )
        self.assertEqual(res.status_code, 400)

    def test_add_profile_keywords_view_not_found(self):
        url = reverse("add_profile_keywords", kwargs={"profile_id": str(uuid.uuid4())})
        res = self.client.post(
            url,
            data=json.dumps({"keywords": "audit"}),
            content_type="application/json",
        )
        self.assertEqual(res.status_code, 404)

    def test_parse_keywords_file_view_txt_success(self):
        txt_content = b"kickback, bribe\nconsulting fee\toff-book;secret commission"
        uploaded_file = SimpleUploadedFile("terms.txt", txt_content, content_type="text/plain")
        res = self.client.post(
            reverse("parse_keywords_file"),
            {"file": uploaded_file},
        )
        self.assertEqual(res.status_code, 200)
        data = res.json()
        self.assertEqual(data["status"], "success")
        self.assertIn("kickback", data["keywords"])
        self.assertIn("bribe", data["keywords"])
        self.assertIn("consulting fee", data["keywords"])
        self.assertIn("off-book", data["keywords"])
        self.assertIn("secret commission", data["keywords"])

    def test_parse_keywords_file_view_xlsx_success(self):
        import io

        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Surveillance Keyword", "Notes"])
        ws.append(["shell company", "priority 1"])
        ws.append(["hawala transfer", "priority 2"])
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)

        uploaded_file = SimpleUploadedFile(
            "surveillance.xlsx",
            buf.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        res = self.client.post(
            reverse("parse_keywords_file"),
            {"file": uploaded_file},
        )
        self.assertEqual(res.status_code, 200)
        data = res.json()
        self.assertEqual(data["status"], "success")
        self.assertIn("shell company", data["keywords"])
        self.assertIn("hawala transfer", data["keywords"])

    def test_parse_keywords_file_view_no_file(self):
        res = self.client.post(reverse("parse_keywords_file"), {})
        self.assertEqual(res.status_code, 400)
        self.assertIn("No file uploaded", res.json()["message"])

    def test_parse_keywords_file_view_invalid_extension(self):
        uploaded_file = SimpleUploadedFile(
            "malicious.exe", b"binary", content_type="application/octet-stream"
        )
        res = self.client.post(
            reverse("parse_keywords_file"),
            {"file": uploaded_file},
        )
        self.assertEqual(res.status_code, 400)
        self.assertIn("Unsupported file type", res.json()["message"])

    def test_upload_profile_keywords_file_view_success(self):
        txt_content = b"unauthorized payment, phantom vendor"
        uploaded_file = SimpleUploadedFile("keywords.txt", txt_content, content_type="text/plain")
        url = reverse("upload_profile_keywords_file", kwargs={"profile_id": str(self.profile.id)})
        res = self.client.post(
            url,
            {"file": uploaded_file},
        )
        self.assertEqual(res.status_code, 200)
        data = res.json()
        self.assertEqual(data["status"], "success")
        self.assertIn("unauthorized payment", data["profile"]["keywords"])
        self.assertIn("phantom vendor", data["profile"]["keywords"])

        # Check DB updated
        self.profile.refresh_from_db()
        self.assertIn("unauthorized payment", self.profile.keywords)

    def test_upload_profile_keywords_file_view_empty_file(self):
        uploaded_file = SimpleUploadedFile("empty.txt", b"   \n\n  ", content_type="text/plain")
        url = reverse("upload_profile_keywords_file", kwargs={"profile_id": str(self.profile.id)})
        res = self.client.post(
            url,
            {"file": uploaded_file},
        )
        self.assertEqual(res.status_code, 400)
        self.assertIn("No valid keywords found", res.json()["message"])

    def test_upload_profile_keywords_file_view_not_found(self):
        uploaded_file = SimpleUploadedFile("test.txt", b"keyword1", content_type="text/plain")
        url = reverse("upload_profile_keywords_file", kwargs={"profile_id": str(uuid.uuid4())})
        res = self.client.post(
            url,
            {"file": uploaded_file},
        )
        self.assertEqual(res.status_code, 404)

    def test_upload_profile_keywords_csv_non_utf8(self):
        # Create a non-utf8 byte string to trigger latin-1 fallback
        csv_content = "header1,Keyword\nval1,fráud\nval2,bribe,extra".encode("latin-1")
        uploaded_file = SimpleUploadedFile(
            "test_non_utf8.csv", csv_content, content_type="text/csv"
        )
        url = reverse("upload_profile_keywords_file", kwargs={"profile_id": str(self.profile.id)})
        res = self.client.post(url, {"file": uploaded_file})
        self.assertEqual(res.status_code, 200)
        self.assertIn("fráud", res.json()["profile"]["keywords"])
        self.assertIn("bribe", res.json()["profile"]["keywords"])

    def test_upload_profile_keywords_csv(self):
        csv_content = b"header1,Keyword\nval1,fraud\nval2,bribe"
        uploaded_file = SimpleUploadedFile("test.csv", csv_content, content_type="text/csv")
        url = reverse("upload_profile_keywords_file", kwargs={"profile_id": str(self.profile.id)})
        res = self.client.post(url, {"file": uploaded_file})
        self.assertEqual(res.status_code, 200)
        self.assertIn("fraud", res.json()["profile"]["keywords"])
        self.assertIn("bribe", res.json()["profile"]["keywords"])

    def test_sync_global_profiles_from_apps(self):
        # Test the sync behavior for q_verify and q_voice
        # Since these apps are installed, the imports should succeed and the loops will run.
        from q_verify.models import VerificationCase

        from .profiles import sync_all_existing_entities_to_profiles

        VerificationCase.objects.create(
            case_ref="TEST-VER-1",
            case_title="Sync Test Case",
            custodian_name="Test Sync Custodian",
            custodian_email="sync@test.com",
            custodian_department="IT",
        )

        from django.utils import timezone
        from q_voice.models import AudioRecording

        AudioRecording.objects.create(
            source_filename="test_sync.wav",
            custodian_name="Test Voice Sync Custodian",
            risk_score=60,
            call_timestamp=timezone.now(),
        )

        from q_bank.models import AuditedPerson

        AuditedPerson.objects.create(
            full_name="Test Bank Sync Custodian",
            notes="flagged",
        )

        created = sync_all_existing_entities_to_profiles()
        self.assertGreaterEqual(created, 3)


class ProfileKeywordRegistryIntegrationTests(TestCase):
    """
    Validates end-to-end profile surveillance keyword registry consumption across
    all forensic Q-apps (q_bank, q_mail, q_voice, q_chat, q_trail, q_scan).
    """

    def setUp(self):
        self.profile = InvestigationProfile.objects.create(
            full_name="Vikram Sethi",
            department="Procurement",
            keywords=["PROJECT_TITAN_SECRET", "SWISS_ACCOUNT", "COMMISSION_CUT", "SHELL_CORP"],
        )
        self.factory = RequestFactory()

    def test_get_profile_keywords_resolver(self):
        from core.profiles import get_profile_keywords

        # 1. By Profile ID
        kws = get_profile_keywords(profile_id=self.profile.id)
        self.assertEqual(
            kws, ["PROJECT_TITAN_SECRET", "SWISS_ACCOUNT", "COMMISSION_CUT", "SHELL_CORP"]
        )

        # 2. By Custodian Name with (Auditee) suffix
        kws_auditee = get_profile_keywords(custodian_name="Vikram Sethi (Auditee)")
        self.assertEqual(
            kws_auditee, ["PROJECT_TITAN_SECRET", "SWISS_ACCOUNT", "COMMISSION_CUT", "SHELL_CORP"]
        )

        # 3. By Request Session Active Profile
        req = self.factory.get("/")
        req.session = {"active_profile_id": str(self.profile.id)}
        kws_session = get_profile_keywords(request=req)
        self.assertEqual(
            kws_session, ["PROJECT_TITAN_SECRET", "SWISS_ACCOUNT", "COMMISSION_CUT", "SHELL_CORP"]
        )

    def test_q_bank_fuzzy_search_profile_keywords(self):
        from q_bank.models import AuditedPerson, BankAccount, BankTransaction

        person = AuditedPerson.objects.create(full_name="Vikram Sethi")
        account = BankAccount.objects.create(
            person=person,
            account_number="9988776655",
            bank_name="HDFC Bank",
        )
        BankTransaction.objects.create(
            account=account,
            txn_date="2026-03-01T10:00:00Z",
            narration="IMPS PAYMENT FOR PROJECT_TITAN_SECRET ROUTING",
            direction=BankTransaction.Direction.DEBIT,
            debit_amount=Decimal("50000.00"),
        )

        client = Client()
        session = client.session
        session["portal_authenticated"] = True
        session["active_profile_id"] = str(self.profile.id)
        session.save()

        # Call fuzzy search without keywords param; should resolve from profile
        res = client.get(f"/bank/api/fuzzy-search/?person_id={person.id}&threshold=70")
        self.assertEqual(res.status_code, 200)
        data = res.json()
        self.assertGreaterEqual(data["total_matches"], 1)

    def test_q_mail_checkpoint_evaluation_profile_keywords(self):
        from q_mail.backend.checkpoints import evaluate_email_checkpoints

        class MockEmail:
            subject = "Confidential: Payment and PROJECT_TITAN_SECRET confirmation"
            body_plain = "Please ensure the shell_corp documentation is attached."
            sender_email = "vikram@external.com"
            recipients_to = ["auditor@hmil.net"]
            recipients_cc = []
            recipients_bcc = []

        res = evaluate_email_checkpoints(MockEmail(), profile_keywords=self.profile.keywords)
        self.assertIn("PROJECT_TITAN_SECRET", res["matched_profile_keywords"])
        self.assertIn("SHELL_CORP", res["matched_profile_keywords"])
        badges = [b["label"] for b in res["badges"]]
        self.assertIn("Profile: PROJECT_TITAN_SECRET", badges)
        self.assertIn("Profile: SHELL_CORP", badges)

    def test_q_voice_tagging_and_screening_profile_keywords(self):
        from q_voice.backend.voice_parser import (
            screen_text_for_intent,
            tag_transcript_detections,
        )

        text = "He mentioned using project_titan_secret to dispatch the parcel."
        dets = tag_transcript_detections(text, extra_keywords=self.profile.keywords)
        types = [d["type"] for d in dets]
        self.assertIn("profile", types)
        terms = [d["term"] for d in dets]
        self.assertIn("PROJECT_TITAN_SECRET", terms)

        intent, flagged, risk = screen_text_for_intent(text, extra_keywords=self.profile.keywords)
        self.assertIn("PROJECT_TITAN_SECRET", flagged)
        self.assertGreaterEqual(risk, 35)

    def test_q_chat_screening_profile_keywords(self):
        from q_chat.backend.chat_parser import screen_message_text

        text = "Meet me near the hotel, bring the project_titan_secret voucher slip."
        risk, flagged = screen_message_text(text, extra_keywords=self.profile.keywords)
        self.assertIn("PROJECT_TITAN_SECRET", flagged)
        self.assertGreaterEqual(risk, 25)

    def test_q_trail_keyword_matching(self):
        from q_trail.services import analyze_profiles_money_trail

        profile2 = InvestigationProfile.objects.create(
            full_name="Rajesh Sharma",
            department="Vendor Management",
            keywords=["OFFSHORE_ACC"],
        )

        res = analyze_profiles_money_trail(profile_ids=[str(self.profile.id), str(profile2.id)])
        self.assertIn("trail_keywords", res)
        self.assertIn("PROJECT_TITAN_SECRET", res["trail_keywords"])
        self.assertIn("OFFSHORE_ACC", res["trail_keywords"])

    def test_q_scan_config_profile_keywords_injection(self):
        client = Client()
        session = client.session
        session["portal_authenticated"] = True
        session["active_profile_id"] = str(self.profile.id)
        session.save()

        res = client.get("/scan/download/config.json/")
        self.assertEqual(res.status_code, 200)
        cfg_data = json.loads(res.content.decode("utf-8"))
        self.assertIn("PROJECT_TITAN_SECRET", cfg_data["keywords"])
        self.assertIn("SWISS_ACCOUNT", cfg_data["keywords"])


class CoreFuzzyEngineTests(TestCase):
    def test_extract_keywords_from_string(self):
        from core.fuzzy import extract_keywords_from_string

        self.assertEqual(extract_keywords_from_string(""), [])
        res = extract_keywords_from_string("alpha, beta; gamma\r\nALPHA, delta")
        self.assertEqual(res, ["alpha", "beta", "gamma", "delta"])

    def test_extract_keywords_from_txt(self):
        from core.fuzzy import extract_keywords_from_txt

        raw_bytes = b"# Comment line\n// Another comment\napple, banana\n\ncherry"
        res = extract_keywords_from_txt(raw_bytes)
        self.assertEqual(res, ["apple", "banana", "cherry"])

        # String input
        res_str = extract_keywords_from_txt("dog, cat\n# ignore\nfish")
        self.assertEqual(res_str, ["dog", "cat", "fish"])

        # Latin-1 encoded bytes
        latin1_bytes = "café, naïve".encode("latin-1")
        res_latin = extract_keywords_from_txt(latin1_bytes)
        self.assertIn("café", res_latin)

    def test_extract_keywords_from_xlsx_general_table(self):
        import io

        import openpyxl

        from core.fuzzy import extract_keywords_from_xlsx

        wb = openpyxl.Workbook()
        ws = wb.active
        # Header with stopword
        ws.append(["keyword", "notes"])
        ws.append(["shell_corp", "internal alert"])
        ws.append(["bribe_fund", "cash"])
        ws.append([None, None])
        buf = io.BytesIO()
        wb.save(buf)

        res = extract_keywords_from_xlsx(buf.getvalue())
        self.assertIn("shell_corp", res)
        self.assertIn("bribe_fund", res)

    def test_extract_keywords_from_xlsx_targeted_header(self):
        import io

        import openpyxl

        from core.fuzzy import extract_keywords_from_xlsx

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Target Name", "Watchlist Identifier", "Date"])
        ws.append(["Vikram", "ACC-9988; HAWALA-1", "2026-01-01"])
        buf = io.BytesIO()
        wb.save(buf)

        res = extract_keywords_from_xlsx(buf.getvalue())
        self.assertIn("ACC-9988", res)
        self.assertIn("HAWALA-1", res)

    def test_extract_keywords_from_xlsx_fallback_no_target_col(self):
        import io

        import openpyxl

        from core.fuzzy import extract_keywords_from_xlsx

        wb = openpyxl.Workbook()
        ws = wb.active
        # Header without any target column words
        ws.append(["Category", "Details"])
        ws.append(["Finance", "hawala_node, shell_account"])
        buf = io.BytesIO()
        wb.save(buf)

        res = extract_keywords_from_xlsx(buf.getvalue())
        self.assertIn("hawala_node", res)
        self.assertIn("shell_account", res)

    def test_extract_keywords_from_file_and_fallbacks(self):
        from core.fuzzy import extract_keywords_from_file

        # TXT file obj
        txt_file = SimpleUploadedFile("watch.csv", b"wire_transfer, kickback")
        res_txt = extract_keywords_from_file(txt_file, "watch.csv")
        self.assertIn("wire_transfer", res_txt)

        # Fallback unknown extension
        unknown_file = SimpleUploadedFile("unknown.dat", b"secret_fund, offshore")
        res_unknown = extract_keywords_from_file(unknown_file, "unknown.dat")
        self.assertIn("secret_fund", res_unknown)

    def test_extract_keywords_from_request(self):
        from core.fuzzy import extract_keywords_from_request

        factory = RequestFactory()

        # 1. From GET query
        req_get = factory.get("/?keywords=alpha,beta")
        res_get = extract_keywords_from_request(req_get)
        self.assertEqual(res_get, ["alpha", "beta"])

        # 2. From POST q
        req_post = factory.post("/", {"q": "gamma; delta"})
        res_post = extract_keywords_from_request(req_post)
        self.assertEqual(res_post, ["gamma", "delta"])

        # 3. From attached file
        f = SimpleUploadedFile("terms.txt", b"epsilon, zeta")
        req_file = factory.post("/", {"file": f})
        res_file = extract_keywords_from_request(req_file)
        self.assertIn("epsilon", res_file)
        self.assertIn("zeta", res_file)

    def test_score_text_against_keywords(self):
        from core.fuzzy import score_text_against_keywords

        # Empty cases
        self.assertEqual(score_text_against_keywords("", ["abc"]), (False, 0, ""))
        self.assertEqual(score_text_against_keywords("some text", []), (False, 0, ""))

        # Direct 100% substring match
        matched, score, kw = score_text_against_keywords(
            "Payment to Apex Logistics", ["Apex Logistics"]
        )
        self.assertTrue(matched)
        self.assertEqual(score, 100)
        self.assertEqual(kw, "Apex Logistics")

        # Fuzzy token match
        matched_fuzz, score_fuzz, kw_fuzz = score_text_against_keywords(
            "Transfer to Logistix", ["Logistics"], threshold=70
        )
        self.assertTrue(matched_fuzz)
        self.assertGreaterEqual(score_fuzz, 70)
        self.assertEqual(kw_fuzz, "Logistics")

        # Below threshold match
        matched_no, _, _ = score_text_against_keywords(
            "Completely unrelated text", ["Logistics"], threshold=90
        )
        self.assertFalse(matched_no)
